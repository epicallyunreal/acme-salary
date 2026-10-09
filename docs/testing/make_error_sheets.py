"""Writes the end-to-end pass's error sheets into docs/testing/errors/ (PLAN step 13).

From backend/: uv run python ../docs/testing/make_error_sheets.py

Each file draws out the messages of SPEC section 7 and section 8 through the Imports page. The
employees a row matches are read from sample-data/base-15000.xlsx, so the sheets fit a database
emptied on the Demo data page and given the base import. The order the pass uses them in, and
the edits it makes in the UI between them, are in docs/TESTING.md.
"""

import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).parent
OUT = HERE / "errors"
BASE = HERE.parent.parent / "sample-data" / "base-15000.xlsx"

COMPONENTS = (
    "BASE",
    "AWS",
    "HOLIDAY_ALLOWANCE",
    "HRA",
    "ALLOWANCES",
    "TAX",
    "CPF",
    "NATIONAL_INSURANCE",
    "SOCIAL_INSURANCE",
    "EMPLOYER_CONTRIBUTION",
)
HEADER = (
    "legacy_code",
    "full_name",
    "email",
    "department",
    "designation",
    "level",
    "country_code",
    "gender",
    "joining_date",
    "last_date",
    "effective_from",
    "bonus_annual",
    *COMPONENTS,
)


def base_rows() -> list[dict]:
    """The base file's rows as dicts, dates as dates."""
    sheet = load_workbook(BASE, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    found = []
    for values in rows:
        row = dict(zip(header, values, strict=True))
        for name in ("joining_date", "last_date", "effective_from"):
            if isinstance(row[name], datetime):
                row[name] = row[name].date()
        found.append(row)
    return found


def staff(rows: list[dict], country: str, skip: int = 0) -> dict:
    """The (skip+1)th employee on staff in `country`, as the base file has them."""
    on_staff = [r for r in rows if r["country_code"] == country and r["last_date"] is None]
    return dict(on_staff[skip])


def leaver(rows: list[dict], country: str) -> dict:
    return dict(next(r for r in rows if r["country_code"] == country and r["last_date"]))


def new_hire(n: int, **changes) -> dict:
    """A valid new employee in India, numbered so no two share a code or an email."""
    row = {
        "legacy_code": f"E2E-{n:03d}",
        "full_name": f"Test Person {n}",
        "email": f"e2e.person.{n}@acme.example",
        "department": "engineering",
        "designation": "Engineer",
        "level": 2,
        "country_code": "IN",
        "gender": "female",
        "joining_date": date(2024, 3, 1),
        "effective_from": date(2024, 3, 1),
        "BASE": 900000,
        "HRA": 360000,
        "ALLOWANCES": 200000,
        "TAX": 100000,
        "EMPLOYER_CONTRIBUTION": 80000,
    }
    row.update(changes)
    return row


def matched(row: dict, **changes) -> dict:
    """A row for an employee on file: identity as it is, amounts as they are, the given changes
    on top. Without an effective_from the row changes identity alone."""
    found = {**row, "effective_from": None}
    found.update(changes)
    return found


def write_xlsx(name: str, rows: list[dict], header=HEADER) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Employees"
    sheet.append(list(header))
    for row in rows:
        sheet.append([row.get(column) for column in header])
    book.save(OUT / name)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    rows = base_rows()

    # Refused before any row is read: one file each.
    (OUT / "not-a-sheet.txt").write_text("legacy_code,full_name\n")
    (OUT / "text-renamed.xlsx").write_text("legacy_code,full_name\nE2E-1,Someone\n")
    book = Workbook()
    book.save(OUT / "empty.xlsx")
    write_xlsx(
        "missing-column.xlsx",
        [new_hire(1)],
        header=[c for c in HEADER if c != "country_code"],
    )
    write_xlsx("column-twice.xlsx", [new_hire(2)], header=[*HEADER, "email"])
    two_codes = ["EmployeeID", "Emp ID", *HEADER[1:]]
    sheet_book = Workbook()
    sheet = sheet_book.active
    sheet.append(two_codes)
    row = new_hire(3)
    sheet.append(["E2E-003", "E2E-003", *[row.get(c) for c in HEADER[1:]]])
    sheet_book.save(OUT / "two-code-columns.xlsx")
    with (OUT / "latin1.csv").open("w", encoding="latin-1", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADER)
        row = new_hire(4, full_name="Zoë Müller")
        writer.writerow([row.get(c, "") for c in HEADER])

    # Every row rule of section 7 that a staged row can show, one row each.
    harini = staff(rows, "IN")
    victor = leaver(rows, "IN")
    too_long = "x" * 300
    row_rules = [
        new_hire(10, legacy_code=None),
        new_hire(11, designation=None),
        new_hire(12, legacy_code="E2E-" + "9" * 29),
        new_hire(13, full_name="N" * 121),
        new_hire(14, designation=too_long),
        new_hire(15, legacy_code="E2E-DUP"),
        new_hire(16, legacy_code="E2E-DUP"),
        new_hire(17, email="e2e.twice@acme.example"),
        new_hire(18, email="e2e.twice@acme.example"),
        new_hire(19, email="not-an-email"),
        new_hire(20, email=harini["email"]),
        new_hire(21, country_code="XX"),
        matched(harini, country_code="US"),
        new_hire(22, department="Growth"),
        new_hire(23, level=9),
        new_hire(24, gender="robot"),
        new_hire(25, joining_date="31/02/2024", effective_from=None),
        new_hire(26, last_date="9999-12-31"),
        new_hire(27, last_date=date(2024, 1, 1)),
        matched(staff(rows, "GB"), joining_date=date(2016, 1, 1)),
        matched(victor, last_date=None),
        new_hire(28, joining_date=date(2024, 3, 1), effective_from=date(2024, 1, 1)),
        matched(
            staff(rows, "SG"),
            BASE=staff(rows, "SG")["BASE"] + 1000,
            effective_from=staff(rows, "SG")["joining_date"],
        ),
        matched(staff(rows, "US"), BASE=staff(rows, "US")["BASE"] + 1000),
        new_hire(
            29,
            country_code="US",
            HRA=100,
            ALLOWANCES=None,
            BASE=90000,
            TAX=9000,
            EMPLOYER_CONTRIBUTION=7000,
        ),
        new_hire(30, BASE=-5000),
        new_hire(31, BASE="1.200.000,50"),
        new_hire(32, BASE="100000.123"),
        new_hire(33, BASE=None),
        new_hire(34, BASE=5, HRA=None, ALLOWANCES=None, TAX=None, EMPLOYER_CONTRIBUTION=None),
        new_hire(35, HRA=12000000000),
        new_hire(36, bonus_annual=-1),
        # Warnings and an info, which leave the row valid.
        new_hire(37, TAX=1500000, BASE=600000, HRA=None, ALLOWANCES=None),
        new_hire(38, bonus_annual=1500000, BASE=1000000),
        new_hire(39, BASE=250000, HRA=50000, ALLOWANCES=None, TAX=10000),
        new_hire(40, BASE=60000000),
        matched(
            staff(rows, "NL"),
            BASE=staff(rows, "NL")["BASE"] - 20000,
            effective_from=date(2026, 11, 1),
        ),
        matched(staff(rows, "DE")),
        new_hire(41, effective_from=None),
    ]
    header = [*HEADER, "Remarks"]
    write_xlsx("row-rules.xlsx", row_rules, header=header)

    # A revision on an employee who has left, alone in its sheet.
    write_xlsx(
        "exited-revision.xlsx",
        [matched(victor, BASE=victor["BASE"] + 1000, effective_from=date(2026, 11, 1))],
    )
    # Two sheets creating the same new code: approve one, then the other (D24).
    write_xlsx("legacy-taken-a.xlsx", [new_hire(50, legacy_code="E2E-TAKEN")])
    write_xlsx(
        "legacy-taken-b.xlsx",
        [new_hire(51, legacy_code="E2E-TAKEN", email="e2e.taken.b@acme.example")],
    )
    # An identity change on an employee whose pay is then revised in the UI, effective today,
    # before the approve, so the row would apply as a revision (D57). A scheduled revision would
    # not do, since it leaves the current record as it was.
    # The row carries an effective_from, so with the pay changed it is still valid and would apply
    # as a revision, rather than fail for want of a date.
    changed = staff(rows, "IN", skip=3)
    write_xlsx(
        "changed-since-upload.xlsx",
        [matched(changed, designation="Lead Engineer", effective_from=date(2026, 12, 1))],
    )
    # An exit for an employee given a scheduled revision in the UI first.
    scheduled = staff(rows, "DE", skip=1)
    write_xlsx("exit-with-scheduled.xlsx", [matched(scheduled, last_date=date(2026, 12, 31))])

    for path in sorted(OUT.iterdir()):
        print(path.name)
    print(f"changed-since-upload: {changed['legacy_code']} {changed['full_name']}")
    print(f"exit-with-scheduled: {scheduled['legacy_code']} {scheduled['full_name']}")


if __name__ == "__main__":
    main()
