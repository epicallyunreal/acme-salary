"""The /admin/ routes of the Demo data page over HTTP (SPEC section 8, F9, D72)."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.demo.views import ResetDemoView
from apps.employees.models import Employee
from apps.imports import services as imports
from apps.imports.parser import XLSX_TYPE
from tests.demo.test_services import PASSWORD, seed

pytestmark = pytest.mark.django_db

RESET = "/api/v1/admin/reset-demo/"
FILES = "/api/v1/admin/sample-files/"
COUNT_KEYS = [
    "employees",
    "compensation_records",
    "compensation_lines",
    "import_batches",
    "import_rows",
    "audit_entries",
]


@pytest.fixture
def seeded(lookups, rates):
    seed(5, 2)


def reset(api, password=PASSWORD, confirm="RESET", **headers):
    return api.post(RESET, {"password": password, "confirm": confirm}, format="json", **headers)


def test_the_reset_is_throttled_on_its_own_scope_five_an_hour(settings):
    """The POST alone: the counts the dialog reads are never throttled."""
    assert ResetDemoView.throttle_scope == "reset"
    assert settings.REST_FRAMEWORK["DEFAULT_THROTTLE_RATES"]["reset"] == "5/hour"


def test_every_route_needs_a_login(anon):
    assert anon.get(RESET).status_code == 401
    assert anon.post(RESET, {"password": "x", "confirm": "RESET"}, format="json").status_code == 401
    assert anon.get(FILES).status_code == 401
    assert anon.get(f"{FILES}base/").status_code == 401


def test_the_counts_come_back_in_a_fixed_number_of_queries(api, seeded, django_assert_num_queries):
    """The user from the token, then one count per table."""
    with django_assert_num_queries(7):
        response = api.get(RESET)

    assert response.status_code == 200
    assert list(response.json()) == COUNT_KEYS
    assert response.json()["employees"] == 7


def test_a_reset_needs_the_password_and_the_word(api, seeded):
    """Wrong or blank, either field is the service's 403 naming it (SPEC section 8); only a
    field left out of the body is a 400."""
    both = reset(api, password="nope", confirm="reset")
    password = reset(api, password="nope")
    confirm = reset(api, confirm="Reset")
    blank = reset(api, password="", confirm="")
    missing = api.post(RESET, {}, format="json")

    assert (both.status_code, both.json()) == (
        403,
        {"password": ["password did not match"], "confirm": ["type RESET to confirm"]},
    )
    assert (password.status_code, password.json()) == (
        403,
        {"password": ["password did not match"]},
    )
    assert (confirm.status_code, confirm.json()) == (403, {"confirm": ["type RESET to confirm"]})
    assert (blank.status_code, list(blank.json())) == (403, ["password", "confirm"])
    assert (missing.status_code, missing.json()) == (
        400,
        {"password": ["This field is required."], "confirm": ["This field is required."]},
    )
    assert Employee.objects.count() == 7


def test_a_reset_returns_the_counts_and_the_user_can_still_log_in(api, anon, seeded, hr_user):
    response = reset(api)

    assert response.status_code == 200
    assert list(response.json()) == COUNT_KEYS and response.json()["employees"] == 7
    assert Employee.objects.count() == 0
    login = anon.post(
        "/api/v1/auth/login/", {"email": hr_user.email, "password": PASSWORD}, format="json"
    )
    assert login.status_code == 200


def test_the_sixth_reset_attempt_within_an_hour_is_429_and_the_counts_are_not(api, seeded):
    for _ in range(5):
        assert reset(api, password="nope").status_code == 403

    sixth = reset(api)

    assert sixth.status_code == 429 and "Retry-After" in sixth
    assert api.get(RESET).status_code == 200
    assert Employee.objects.count() == 7


def test_the_reset_throttle_counts_per_client_address_not_per_login(api, seeded):
    """One reviewer's five attempts never block another on the shared login (D72)."""
    first = {"HTTP_X_FORWARDED_FOR": "203.0.113.5"}
    second = {"HTTP_X_FORWARDED_FOR": "203.0.113.9"}
    for _ in range(5):
        assert reset(api, password="nope", **first).status_code == 403

    assert reset(api, password="nope", **first).status_code == 429
    assert reset(api, password="nope", **second).status_code == 403


def test_a_reset_while_a_file_is_being_read_is_409(api, seeded):
    assert imports._reading.acquire(blocking=False)
    try:
        response = reset(api)
    finally:
        imports._reading.release()

    assert (response.status_code, response.json()) == (
        409,
        {"detail": "an import is being read; try again in a moment"},
    )


def test_the_file_list_and_the_downloads(api, lookups):
    listed = api.get(FILES)
    base = api.get(f"{FILES}base/")
    generated = api.get(f"{FILES}exits/")
    unknown = api.get(f"{FILES}nope/")
    bad_seed = api.get(f"{FILES}base/?seed=x")

    assert listed.status_code == 200 and listed.json()["employees"] == 0
    assert [f["name"] for f in listed.json()["files"]][:2] == ["base", "newhires"]
    assert base.status_code == 200 and base["Content-Type"] == XLSX_TYPE
    assert base["Content-Disposition"] == 'attachment; filename="base-15000.xlsx"'
    assert (generated.status_code, generated.json()) == (
        409,
        {"detail": "import the base file first"},
    )
    assert unknown.status_code == 404
    assert (bad_seed.status_code, bad_seed.json()) == (
        400,
        {"seed": ["seed must be a whole number"]},
    )


def rows_of(response) -> list[tuple]:
    sheet = load_workbook(BytesIO(response.content), read_only=True).worksheets[0]
    return list(sheet.iter_rows(values_only=True))


def test_a_generated_file_downloads_and_another_seed_draws_differently(api, lookups, rates):
    """`?seed=` steers the draws: diwali-bonus's shares of base differ between 42 and 7."""
    seed(20, 5)

    response = api.get(f"{FILES}diwali-bonus/")
    other = api.get(f"{FILES}diwali-bonus/?seed=7")

    assert response.status_code == 200 and response["Content-Type"] == XLSX_TYPE
    assert response["Content-Disposition"] == 'attachment; filename="diwali-bonus.xlsx"'
    header, *rows = rows_of(response)
    assert header[:3] == ("legacy_code", "full_name", "email")
    assert rows and rows != rows_of(other)[1:]
    assert [row[0] for row in rows] == [row[0] for row in rows_of(other)[1:]]


def test_the_file_list_and_a_generated_download_run_a_fixed_number_of_queries(
    api, seeded, django_assert_num_queries
):
    """The user from the token and the employee count for the list. For a generated file, the
    count, the share's employees with their country and department, their open records for the
    exit rules, their latest and current records and the lines, and the component codes."""
    with django_assert_num_queries(2):
        listed = api.get(FILES)
    with django_assert_num_queries(8):
        exits = api.get(f"{FILES}exits/")

    assert listed.status_code == 200 and exits.status_code == 200
