"""The demo reset and the sample files (SPEC F9, D72), without HTTP."""

from datetime import date, timedelta
from io import BytesIO
from itertools import combinations
from pathlib import Path

import pytest
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import OperationalError
from django.http import Http404
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.compensation.models import CompensationLine, CompensationRecord
from apps.compensation.services import bulk_apply
from apps.compensation.validators import template_columns
from apps.core.exceptions import Conflict, Forbidden
from apps.core.models import Country, Department, ExchangeRate, SalaryComponent
from apps.demo import services
from apps.employees.models import Employee
from apps.imports import services as imports
from apps.imports.applier import approve
from apps.imports.management.commands.make_import_file import database_codes
from apps.imports.models import ImportBatch, ImportRow
from apps.imports.parser import XLSX_TYPE
from seed import profiles
from seed.generator import ordered_entries
from tests.imports.sheets import new_row, upload_of

pytestmark = pytest.mark.django_db

PASSWORD = "correct-horse-battery"  # UserFactory's
TODAY = date(2026, 10, 7)


def seed(active: int, exited: int, seed_value: int = 42) -> None:
    bulk_apply(
        ordered_entries(active, seed_value, exited=exited),
        source="seed",
        at=profiles.REFERENCE_DATETIME,
        today=profiles.REFERENCE_DATE,
    )


@pytest.fixture
def seeded(lookups, rates):
    """Five on staff and two leavers from the seed, and one staged import."""
    seed(5, 2)
    imports.stage_upload(upload_of(new_row()), actor=None)


def lookup_counts() -> tuple[int, int, int, int]:
    return (
        Country.objects.count(),
        Department.objects.count(),
        SalaryComponent.objects.count(),
        ExchangeRate.objects.count(),
    )


def test_a_reset_empties_the_pay_tables_keeps_the_rest_and_restarts_the_ids(seeded, hr_user):
    before = services.counts()
    assert before["employees"] == 7 and before["import_batches"] == 1
    assert before["audit_entries"] > 0 and before["import_rows"] == 1
    lookups_before = lookup_counts()

    deleted = services.reset_demo(actor=hr_user, password=PASSWORD, confirm="RESET")

    assert deleted == before
    assert services.counts() == {**dict.fromkeys(before, 0), "audit_entries": 1}
    assert lookup_counts() == lookups_before
    assert get_user_model().objects.get(pk=hr_user.pk).check_password(PASSWORD)
    (entry,) = AuditLog.objects.all()
    assert (entry.action, entry.actor_id, entry.employee_id) == ("demo_reset", hr_user.pk, None)
    assert (entry.entity_type, entry.entity_id, entry.before, entry.after) == (
        "demo_database",
        0,
        before,
        {},
    )
    # The sequence counters are cleared, so the next employee is ACME-1 and its record id 1.
    seed(1, 0, seed_value=7)
    assert Employee.objects.get().pk == 1
    assert CompensationRecord.objects.order_by("pk").first().pk == 1
    assert CompensationLine.objects.order_by("pk").first().pk == 1
    assert ImportBatch.objects.count() == 0 and ImportRow.objects.count() == 0


def test_a_wrong_password_or_confirmation_is_refused_naming_each_wrong_field(seeded, hr_user):
    with pytest.raises(Forbidden) as both:
        services.reset_demo(actor=hr_user, password="nope", confirm="reset")
    with pytest.raises(Forbidden) as password:
        services.reset_demo(actor=hr_user, password="nope", confirm="RESET")
    with pytest.raises(Forbidden) as confirm:
        services.reset_demo(actor=hr_user, password=PASSWORD, confirm="Reset")

    assert both.value.errors == {
        "password": ["password did not match"],
        "confirm": ["type RESET to confirm"],
    }
    assert password.value.errors == {"password": ["password did not match"]}
    assert confirm.value.errors == {"confirm": ["type RESET to confirm"]}
    assert Employee.objects.count() == 7
    assert not AuditLog.objects.filter(action="demo_reset").exists()


def test_a_reset_while_a_file_is_being_read_or_applied_is_a_conflict(seeded, hr_user, monkeypatch):
    """The import lock covers a file being read. SQLite's write lock covers an approve in
    flight: the transaction waits out the 20 s and the locked error becomes a 409, the truncation
    already done rolled back with it, and the lock is free again afterwards (D72). Any other
    database error is the 500 it always was."""
    assert imports._reading.acquire(blocking=False)
    try:
        with pytest.raises(Conflict, match="an import is being read; try again in a moment"):
            services.reset_demo(actor=hr_user, password=PASSWORD, confirm="RESET")
    finally:
        imports._reading.release()
    before = services.counts()

    def locked(*args, **kwargs):  # the audit entry is written after the truncation
        raise OperationalError("database is locked")

    monkeypatch.setattr(services.audit, "record", locked)
    with pytest.raises(Conflict, match="an import is being applied; try again in a moment"):
        services.reset_demo(actor=hr_user, password=PASSWORD, confirm="RESET")
    assert services.counts() == before

    def broken(*args, **kwargs):
        raise OperationalError("no such table: audit_auditlog")

    monkeypatch.setattr(services.audit, "record", broken)
    with pytest.raises(OperationalError, match="no such table"):
        services.reset_demo(actor=hr_user, password=PASSWORD, confirm="RESET")
    assert services.counts() == before
    assert Employee.objects.count() == 7

    monkeypatch.undo()
    assert services.reset_demo(actor=hr_user, password=PASSWORD, confirm="RESET") == before


def test_the_file_list_and_the_base_file_need_no_data(lookups):
    payload = services.sample_files_payload()

    assert payload["employees"] == 0
    assert [(f["name"], f["available"]) for f in payload["files"]] == [
        ("base", True),
        ("newhires", False),
        ("diwali-bonus", False),
        ("april-raises", False),
        ("promotions", False),
        ("exits", False),
        ("fix-me", False),
    ]
    assert all(f["description"] and f["file_name"].endswith(".xlsx") for f in payload["files"])
    name, data = services.sample_file("base")
    assert name == "base-15000.xlsx"
    assert data == (Path(settings.SAMPLE_DATA_DIR) / "base-15000.xlsx").read_bytes()
    with pytest.raises(Conflict, match="import the base file first"):
        services.sample_file("exits")
    with pytest.raises(Http404):
        services.sample_file("nope")


def rows_of(data: bytes) -> tuple[list, list[dict]]:
    values = list(
        load_workbook(BytesIO(data), read_only=True).worksheets[0].iter_rows(values_only=True)
    )
    header = list(values[0])
    return header, [dict(zip(header, row, strict=True)) for row in values[1:]]


def test_generated_files_come_from_the_live_tables_in_the_templates_shape(seeded):
    """Seven employees on 7 October 2026: exits takes those on staff of share 3 whose drawn last
    day, 14 days before to 60 after, they can leave on, fix-me the raises of share 4 with 20 new
    joiners and its planted errors, newhires 500 who joined between 60 days ago and 14 ahead."""
    assert all(f["available"] for f in services.sample_files_payload()["files"])
    columns = template_columns(database_codes())
    on_staff = {e.pk: e for e in Employee.objects.filter(last_date=None)}

    exits_name, exits = services.sample_file("exits", today=TODAY)
    header, rows = rows_of(exits)
    assert exits_name == "exits.xlsx" and header == columns
    assert [Employee.objects.get(legacy_code=r["legacy_code"]).pk % 5 for r in rows] == [3] * len(
        rows
    )
    assert all(
        TODAY - timedelta(days=14) <= r["last_date"].date() <= TODAY + timedelta(days=60)
        for r in rows
    )
    assert len(rows) <= sum(pk % 5 == 3 for pk in on_staff)

    _, fixme = services.sample_file("fix-me", today=TODAY)
    header, rows = rows_of(fixme)
    assert header == columns
    assert len(rows) == sum(pk % 5 == 4 for pk in on_staff) + 20
    assert rows[0]["BASE"] == -5000

    _, hires = services.sample_file("newhires", today=TODAY)
    _, rows = rows_of(hires)
    assert len(rows) == 500
    joined = [r["joining_date"].date() for r in rows]
    assert all(date(2026, 8, 8) <= day <= date(2026, 10, 21) for day in joined)
    assert min(joined) < TODAY < max(joined)
    assert len({r["legacy_code"] for r in rows}) == 500


GENERATED = ("newhires", "diwali-bonus", "april-raises", "promotions", "exits", "fix-me")
SHARES = {"diwali-bonus": 0, "april-raises": 1, "promotions": 2, "exits": 3}


def upload(file_name: str, data: bytes) -> SimpleUploadedFile:
    return SimpleUploadedFile(file_name, data, content_type=XLSX_TYPE)


def test_the_six_files_from_one_database_touch_no_one_twice_and_import_in_turn(lookups, rates):
    """D72 on forty on staff and ten leavers: made from one database, no two files share a
    legacy code or an email, each revision file keeps to its share, april-raises is dated 1 April
    next, and in page order every file imports with no error rows but fix-me's planted ones,
    which mark seven rows since the duplicated code marks both. Imported again, every row that
    is not one of those is a no-op."""
    seed(40, 10)
    on_staff = list(Employee.objects.filter(last_date=None).values_list("pk", flat=True))
    made = {name: services.sample_file(name, today=TODAY) for name in GENERATED}
    rows = {name: rows_of(data)[1] for name, (_, data) in made.items()}
    codes = {name: {r["legacy_code"] for r in rows[name]} for name in GENERATED}
    emails = {name: {r["email"] for r in rows[name]} for name in GENERATED}

    for first, second in combinations(GENERATED, 2):
        assert not codes[first] & codes[second], (first, second)
        assert not emails[first] & emails[second], (first, second)
    for name, share in SHARES.items():
        picked = [Employee.objects.get(legacy_code=r["legacy_code"]).pk for r in rows[name]]
        assert picked and all(pk % 5 == share for pk in picked), name
    assert all(r["effective_from"].date() == date(2027, 4, 1) for r in rows["april-raises"])
    assert all(r["effective_from"].date() == TODAY for r in rows["diwali-bonus"])
    last_days = [r["last_date"].date() for r in rows["exits"]]
    assert all(TODAY - timedelta(days=14) <= day <= TODAY + timedelta(days=60) for day in last_days)
    assert len(rows["newhires"]) == 500
    assert len(rows["fix-me"]) == sum(pk % 5 == 4 for pk in on_staff) + 20

    for name, (file_name, data) in made.items():
        batch = imports.stage_upload(upload(file_name, data), actor=None)
        planted = 7 if name == "fix-me" else 0
        assert (batch.total_rows, batch.error_rows) == (len(rows[name]), planted), name
        for row in batch.rows.filter(status=ImportRow.Status.ERROR):
            imports.skip_row(batch, row)
        approve(batch, actor=None)
    assert Employee.objects.count() == 50 + 500 + 20
    assert Employee.objects.exclude(last_date=None).count() == 10 + len(rows["exits"])
    # A past last day has left the counts, a future one is on staff until the day, and a future
    # joiner is Upcoming until theirs.
    past, future = [d for d in last_days if d < TODAY], [d for d in last_days if d > TODAY]
    assert Employee.objects.filter(last_date__lt=TODAY).count() == 10 + len(past)
    assert Employee.objects.filter(last_date__gt=TODAY).count() == len(future)
    assert Employee.objects.filter(joining_date__gt=TODAY).exists()

    for name, (file_name, data) in made.items():
        again = imports.stage_upload(upload(file_name, data), actor=None)
        assert again.noop_rows == again.total_rows - again.error_rows, name
