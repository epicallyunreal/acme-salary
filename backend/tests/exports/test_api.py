"""GET /api/v1/exports/employees/ and /exports/insights/<name>/ over HTTP (SPEC F6, D60).

The people are DESIGN section 9's thirty, and every expected figure is a literal worked out from
that table; the insight ones are the same figures tests/insights asserts on the cards.
"""

import csv
import io
import json
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from django.test import override_settings
from openpyxl import load_workbook

from apps.exports import tables
from apps.insights import queries
from tests.factories import CompensationRecordFactory, EmployeeFactory, country, department
from tests.insights.fixture import build_fixture, in_lines

pytestmark = pytest.mark.django_db

EMPLOYEES = "/api/v1/exports/employees/"
INSIGHT = "/api/v1/exports/insights/{}/"
RATES = "October 2026 rates"
HEADER = [
    "Employee code",
    "Legacy code",
    "Full name",
    "Email",
    "Department",
    "Designation",
    "Level",
    "Country",
    "Currency",
    "Gross (local)",
    "Bonus (local)",
    "Gross (USD)",
    "Status",
    "Joining date",
    "Last date",
    "Effective from",
    "Version",
]


@pytest.fixture
def people(lookups):
    return build_fixture()


@pytest.fixture
def asha(people):
    """One more Indian engineer, with a legacy code and a bonus, for a row read cell by cell."""
    employee = EmployeeFactory(
        full_name="Asha Rao",
        email="asha.rao@acme.example",
        legacy_code="EMP-0042",
        department=department("engineering"),
        country=country("IN"),
        level=3,
    )
    CompensationRecordFactory(
        employee=employee,
        lines=in_lines(Decimal(1_800_000)),
        bonus_annual=Decimal(200_000),
        effective_from=date(2024, 1, 1),
    )
    return employee


def content(response) -> bytes:
    return b"".join(response.streaming_content) if response.streaming else response.content


def csv_rows(response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content(response).decode("utf-8-sig"))))


def xlsx_rows(response) -> list[list]:
    sheet = load_workbook(io.BytesIO(content(response))).active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def json_body(response) -> dict:
    return json.loads(content(response))


# The directory


@pytest.mark.parametrize("file", ["csv", "xlsx", "json"])
def test_a_filtered_export_has_exactly_the_rows_the_pages_show(api, people, file):
    """SPEC F6: the whole filtered set in the grid's order, whatever page was asked for, on
    Active as the grid opens, so X1 and X2, who left, are in neither (D64)."""
    query = {"country": "IN,US", "level": "1,2,3", "ordering": "-gross"}
    codes = []
    for page in (1, 2, 3):
        body = api.get("/api/v1/employees/", {**query, "page": page, "page_size": 5}).json()
        codes += [row["employee_code"] for row in body["results"]]

    response = api.get(EMPLOYEES, {**query, "file": file, "page": 2, "page_size": 5})

    if file == "csv":
        exported = [row[0] for row in csv_rows(response)[1:]]
    elif file == "xlsx":
        exported = [row[0] for row in xlsx_rows(response)[1:]]
    else:
        exported = [row["employee_code"] for row in json_body(response)["rows"]]
    assert (body["count"], len(codes)) == (12, 12)
    assert exported == codes


def test_without_ordering_the_export_is_in_code_order(api, people):
    """D63: the grid's default sort, Code ascending. DESIGN 9's people were hired U1 to U10
    first, so by name D1 would lead."""
    rows = csv_rows(api.get(EMPLOYEES, {"file": "csv"}))[1:]
    ids = [int(row[0].removeprefix("ACME-")) for row in rows]

    assert ids == sorted(ids)
    assert [row[2] for row in rows[:3]] == ["U1", "U2", "U3"]


def test_the_csv_has_seventeen_columns_in_spec_order(api, asha):
    rows = csv_rows(api.get(EMPLOYEES, {"file": "csv", "search": "asha"}))

    assert rows == [
        HEADER,
        [
            asha.employee_code,
            "EMP-0042",
            "Asha Rao",
            "asha.rao@acme.example",
            "Engineering",
            "Senior Engineer",
            "3",
            "India",
            "INR",
            "2000000.00",
            "200000.00",
            "25000.00",
            "active",
            "2020-01-01",
            "",
            "2024-01-01",
            "1",
        ],
    ]


def test_the_reporting_column_is_headed_with_its_currency(api, asha):
    rows = xlsx_rows(api.get(EMPLOYEES, {"file": "xlsx", "search": "asha", "currency": "eur"}))

    assert rows[0][11] == "Gross (EUR)"
    assert rows[1][9:12] == [2000000, 200000, 20000]


def test_the_json_names_each_value_and_the_filters_it_applied(api, asha):
    """The pages and a parameter the export does not read are not shown as applied (D60)."""
    query = {"search": "asha", "country": "IN", "page": 2, "page_size": 5, "evil": "1"}
    body = json_body(api.get(EMPLOYEES, {**query, "file": "json"}))

    assert body == {
        "generated_at": "2026-10-07T09:00:00+00:00",
        "filters": {
            "search": "asha",
            "country": "IN",
            "on": "2026-10-07",
            "currency": "USD",
            "status": "active",
        },
        "rows": [
            {
                "employee_code": asha.employee_code,
                "legacy_code": "EMP-0042",
                "full_name": "Asha Rao",
                "email": "asha.rao@acme.example",
                "department": "Engineering",
                "designation": "Senior Engineer",
                "level": 3,
                "country": "India",
                "currency": "INR",
                "gross_local": "2000000.00",
                "bonus_local": "200000.00",
                "gross_reporting": "25000.00",
                "status": "active",
                "joining_date": "2020-01-01",
                "last_date": None,
                "effective_from": "2024-01-01",
                "version": 1,
            }
        ],
    }


def test_a_name_starting_with_equals_is_escaped_in_csv_and_xlsx(api, people):
    """SPEC F6: a spreadsheet shows the text instead of running it. JSON keeps it as it is."""
    EmployeeFactory(full_name="=cmd|' /C calc'!A0", email="calc@acme.example")
    query = {"search": "calc"}

    (_, csv_row) = csv_rows(api.get(EMPLOYEES, {**query, "file": "csv"}))
    sheet = load_workbook(io.BytesIO(content(api.get(EMPLOYEES, {**query, "file": "xlsx"}))))
    xlsx_cell = sheet.active["C2"]
    (json_row,) = json_body(api.get(EMPLOYEES, {**query, "file": "json"}))["rows"]

    assert csv_row[2] == "'=cmd|' /C calc'!A0"
    assert (xlsx_cell.value, xlsx_cell.data_type) == ("'=cmd|' /C calc'!A0", "s")
    assert json_row["full_name"] == "=cmd|' /C calc'!A0"


def test_an_employee_with_no_live_record_is_exported_without_amounts(api, people):
    """D50: an exit before pay started leaves the only record cancelled."""
    employee = EmployeeFactory(full_name="Never Paid", email="never@acme.example")
    CompensationRecordFactory(employee=employee, cancelled_at=datetime(2026, 9, 1, tzinfo=UTC))

    (_, row) = csv_rows(api.get(EMPLOYEES, {"file": "csv", "search": "never"}))

    assert row[8:12] == ["INR", "", "", ""]
    assert row[15:] == ["", ""]


def test_the_directory_export_shows_each_record_as_on_its_reporting_date(api, people):
    """On 1 January 2027 U1's scheduled revision to 130,000 is current; U2 has none."""
    query = {"file": "json", "country": "US", "level": "1", "on": "2027-01-01"}
    body = json_body(api.get(EMPLOYEES, query))

    assert body["filters"] == {
        "country": "US",
        "level": "1",
        "on": "2027-01-01",
        "currency": "USD",
        "status": "active",
    }
    assert [
        (row["full_name"], row["gross_local"], row["effective_from"], row["version"])
        for row in body["rows"]
    ] == [("U1", "130000.00", "2027-01-01", 2), ("U2", "110000.00", "2024-01-01", 1)]


def test_a_missing_rate_is_a_400_before_the_csv_starts(api, people):
    """The fixture has no rate for SGD. Found on a row, it would cut the CSV off after its 200."""
    singapore = EmployeeFactory(country=country("SG"), email="sg@acme.example")
    CompensationRecordFactory(employee=singapore, effective_from=date(2024, 1, 1))

    response = api.get(EMPLOYEES, {"file": "csv"})

    assert (response.status_code, response.json()) == (
        400,
        {"currency": ["no exchange rate from SGD to USD"]},
    )


def test_a_row_that_fails_fails_before_the_csv_starts(api, people, monkeypatch):
    """The first rows are read before the 200 goes out, so a failure is an error response."""

    def broken(*args):
        raise RuntimeError("row")

    monkeypatch.setattr(tables, "directory_row", broken)

    with pytest.raises(RuntimeError):
        api.get(EMPLOYEES, {"file": "csv"})


def test_the_default_directory_export_holds_the_active_employees_and_any_holds_everyone(
    api, people
):
    """F6's measured criterion reads the default export (D64): 27 active here, 30 with `any`."""
    active = csv_rows(api.get(EMPLOYEES, {"file": "csv"}))[1:]
    everyone = csv_rows(api.get(EMPLOYEES, {"file": "csv", "status": "any"}))[1:]

    assert (len(active), len(everyone)) == (27, 30)
    assert {row[12] for row in active} == {"active"}


def test_the_directory_export_reads_in_a_fixed_number_of_queries(
    api, people, django_assert_num_queries
):
    """The user, the rates, the currencies of the live records, the employees and their records:
    one prefetch a chunk of 1,000. The 27 active employees and the header: X1 and X2 have left
    and D5 joins in December (D64)."""
    with django_assert_num_queries(5):
        rows = csv_rows(api.get(EMPLOYEES, {"file": "csv"}))

    assert len(rows) == 28


@pytest.mark.parametrize(
    ("file", "content_type"),
    [
        ("csv", "text/csv; charset=utf-8"),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("json", "application/json"),
        ("CSV", "text/csv; charset=utf-8"),
    ],
)
@override_settings(TIME_ZONE="Asia/Kolkata")
def test_the_file_is_an_attachment_stamped_in_utc(api, people, file, content_type):
    """SPEC F6: CSV streams, XLSX and JSON are built whole."""
    response = api.get(EMPLOYEES, {"file": file})

    assert response.status_code == 200
    assert response.streaming is (file.lower() == "csv")
    assert response["Content-Type"] == content_type
    assert response["Content-Disposition"] == (
        f'attachment; filename="acme-employees-20261007-0900.{file.lower()}"'
    )
    assert response["Cache-Control"] == "no-store"


@pytest.mark.parametrize(
    ("params", "errors"),
    [
        ({}, {"file": ["file must be csv, xlsx or json"]}),
        ({"file": "pdf"}, {"file": ["file must be csv, xlsx or json, got pdf"]}),
        (
            {"file": "csv", "currency": "local", "country": "IN,US"},
            {"currency": ["currency=local needs exactly one country in the filter"]},
        ),
        (
            {"file": "csv", "level": "9"},
            {"level": ["Select a valid choice. 9 is not one of the available choices."]},
        ),
        (
            {"file": "csv", "designation": ",".join(["Engineer"] * 21)},
            {"designation": ["designation accepts at most 20 titles"]},
        ),
    ],
)
def test_a_bad_parameter_is_the_directorys_400(api, lookups, params, errors):
    response = api.get(EMPLOYEES, params)

    assert (response.status_code, response.json()) == (400, errors)


# The insights


def test_every_insight_has_an_export_table():
    """An insight added later cannot 404 on its export alone."""
    assert set(tables.INSIGHT_TABLES) == set(queries.INSIGHTS)


@pytest.mark.parametrize("name", sorted(queries.INSIGHTS))
def test_every_insight_table_ends_with_reason_currency_and_rates(api, people, name):
    """SPEC F6 and D60, for the insights promoted after step 7 as well."""
    (header, *_) = csv_rows(api.get(INSIGHT.format(name), {"file": "csv"}))

    assert header[-3:] == ["Reason", "Currency", "Rates"]


def test_the_summary_is_one_row_with_its_currency_and_rates(api, people):
    body = json_body(api.get(INSIGHT.format("summary"), {"file": "json"}))

    assert body["filters"] == {"on": "2026-10-07", "currency": "USD"}
    assert body["rows"] == [
        {
            "currency": "USD",
            "rates": RATES,
            "headcount": 27,
            "total_cost": "2435450.00",
            "mean_gross": "80462.96",
            "median_gross": "75000.00",
            "countries": 4,
            "departments": 4,
            "reason": None,
        }
    ]


def test_a_summary_under_five_says_why_its_figures_are_missing(api, people):
    """D59 and D60: Finance is G1 alone."""
    rows = csv_rows(api.get(INSIGHT.format("summary"), {"file": "csv", "department": "finance"}))

    assert rows == [
        [
            "Headcount",
            "Total cost",
            "Mean gross",
            "Median gross",
            "Countries",
            "Departments",
            "Reason",
            "Currency",
            "Rates",
        ],
        ["1", "", "", "", "1", "1", "fewer than 5 people", "USD", RATES],
    ]


def test_by_country_lists_every_country_with_the_withheld_ones_saying_why(api, people):
    rows = csv_rows(api.get(INSIGHT.format("by-country"), {"file": "csv"}))

    assert rows == [
        [
            "Country code",
            "Country",
            "Headcount",
            "Total cost",
            "Median gross",
            "Share %",
            "Reason",
            "Currency",
            "Rates",
        ],
        ["US", "United States", "10", "1595000.00", "145000.00", "65.5", "", "USD", RATES],
        ["IN", "India", "12", "319200.00", "23750.00", "13.1", "", "USD", RATES],
        ["DE", "Germany", "4", "", "", "", "fewer than 5 people", "USD", RATES],
        ["NL", "Netherlands", "0", "", "", "", "", "USD", RATES],
        ["SG", "Singapore", "0", "", "", "", "", "USD", RATES],
        ["GB", "United Kingdom", "1", "", "", "", "fewer than 5 people", "USD", RATES],
    ]


def test_by_department_shows_its_figures_and_says_why_any_are_missing(api, people):
    """Engineering is the fifteen of tests/insights; Finance is G1 alone and Legal has nobody."""
    query = {"file": "json", "department": "engineering,finance,legal"}
    body = json_body(api.get(INSIGHT.format("by-department"), query))

    assert body["rows"] == [
        {
            "department_code": "engineering",
            "department": "Engineering",
            "headcount": 15,
            "total_cost": "1219200.00",
            "mean_gross": "71500.00",
            "median_gross": "87500.00",
            "reason": None,
            "currency": "USD",
            "rates": RATES,
        },
        {
            "department_code": "finance",
            "department": "Finance",
            "headcount": 1,
            "total_cost": None,
            "mean_gross": None,
            "median_gross": None,
            "reason": "fewer than 5 people",
            "currency": "USD",
            "rates": RATES,
        },
        {
            "department_code": "legal",
            "department": "Legal",
            "headcount": 0,
            "total_cost": None,
            "mean_gross": None,
            "median_gross": None,
            "reason": None,
            "currency": "USD",
            "rates": RATES,
        },
    ]


def test_the_distribution_is_one_row_per_band_numbered_from_one(api, people):
    """The 27 people of tests/insights: p1 10,650.00, p99 187,400.00, twelve bands."""
    rows = csv_rows(api.get(INSIGHT.format("distribution"), {"file": "csv"}))
    edges = ["10650.00", "25379.17", "40108.33", "54837.50", "69566.67", "84295.83"]
    edges += ["99025.00", "113754.17", "128483.33", "143212.50", "157941.67", "172670.83"]
    counts = ["7", "5", "0", "1", "1", "1", "4", "1", "2", "1", "2", "2"]

    assert rows[0] == ["Band", "From", "To", "People", "Reason", "Currency", "Rates"]
    assert rows[1:] == [
        [
            str(band),
            edges[band - 1],
            (edges + ["187400.00"])[band],
            counts[band - 1],
            "",
            "USD",
            RATES,
        ]
        for band in range(1, 13)
    ]


@pytest.mark.parametrize(
    ("query", "people_count", "reason"),
    [({"country": "US"}, 10, "fewer than 20 people"), ({"department": "legal"}, 0, None)],
    ids=["under twenty", "nobody"],
)
def test_a_distribution_without_bands_is_one_row_of_its_headcount(
    api, people, query, people_count, reason
):
    """D60: the file still says how many people the filters matched."""
    body = json_body(api.get(INSIGHT.format("distribution"), {**query, "file": "json"}))

    assert body["rows"] == [
        {
            "band": None,
            "from": None,
            "to": None,
            "people": people_count,
            "reason": reason,
            "currency": "USD",
            "rates": RATES,
        }
    ]


LEVEL_HEADER = [
    "Level",
    "Headcount",
    "p10",
    "p25",
    "Median",
    "p75",
    "p90",
    "Reason",
    "Currency",
    "Rates",
]


def test_a_level_under_twenty_carries_its_headcount_and_empty_figures(api, people):
    """D59: no level of the fixture reaches twenty, so every level is its headcount alone, as the
    card shows it."""
    rows = csv_rows(api.get(INSIGHT.format("by-level"), {"file": "csv"}))

    assert rows == [LEVEL_HEADER] + [
        [str(level), str(count), "", "", "", "", "", "fewer than 20 people", "USD", RATES]
        for level, count in enumerate([4, 6, 5, 5, 4, 3], start=1)
    ]


def test_a_withheld_level_is_null_in_json(api, people):
    rows = json_body(api.get(INSIGHT.format("by-level"), {"file": "json"}))["rows"]

    assert rows == [
        {
            "level": level,
            "headcount": count,
            "p10": None,
            "p25": None,
            "median": None,
            "p75": None,
            "p90": None,
            "reason": "fewer than 20 people",
            "currency": "USD",
            "rates": RATES,
        }
        for level, count in enumerate([4, 6, 5, 5, 4, 3], start=1)
    ]


def test_a_level_of_twenty_exports_p10_to_p90(api, lookups):
    """D59's twenty Dutch salaries of 50,000 to 97,500 USD at level 1, nineteen at level 2: the
    figures of tests/insights, by the inclusive method."""
    build_fixture()
    for step in range(20):
        CompensationRecordFactory(
            employee=EmployeeFactory(country=country("NL"), level=1),
            lines={"BASE": Decimal(40_000 + 2_000 * step)},
            effective_from=date(2024, 1, 1),
        )
    for _ in range(19):
        CompensationRecordFactory(
            employee=EmployeeFactory(country=country("NL"), level=2),
            lines={"BASE": Decimal(60_000)},
            effective_from=date(2024, 1, 1),
        )

    rows = csv_rows(api.get(INSIGHT.format("by-level"), {"file": "csv", "country": "NL"}))

    assert rows[1:3] == [
        ["1", "20", "54750.00", "61875.00", "73750.00", "85625.00", "92750.00", "", "USD", RATES],
        ["2", "19", "", "", "", "", "", "fewer than 20 people", "USD", RATES],
    ]
    assert [row[1] for row in rows[3:]] == ["0", "0", "0", "0"]


def test_the_gender_gap_is_one_row_per_group_and_gender(api, people):
    """India: six women at a median of 2.0M INR, six men at 1.8M, a gap of -11.1. Overall is the
    same group when the filter is India alone."""
    rows = csv_rows(api.get(INSIGHT.format("gender-gap"), {"file": "csv", "country": "IN"}))

    def group(code, name):
        return [
            [code, name, "female", "6", "25000.00", "-11.1", "", "USD", RATES],
            [code, name, "male", "6", "22500.00", "-11.1", "", "USD", RATES],
            [code, name, "other", "0", "", "-11.1", "", "USD", RATES],
            [code, name, "undisclosed", "0", "", "-11.1", "", "USD", RATES],
        ]

    assert rows[0] == [
        "Country code",
        "Country",
        "Gender",
        "Headcount",
        "Median gross",
        "Gap %",
        "Reason",
        "Currency",
        "Rates",
    ]
    assert rows[1:] == group("IN", "India") + group("", "Overall")


def test_the_overall_gender_rows_have_no_country_code(api, people):
    query = {"file": "json", "country": "IN"}
    overall = json_body(api.get(INSIGHT.format("gender-gap"), query))["rows"][4]

    assert overall == {
        "country_code": None,
        "country": "Overall",
        "gender": "female",
        "headcount": 6,
        "median_gross": "25000.00",
        "gap_pct": "-11.1",
        "reason": None,
        "currency": "USD",
        "rates": RATES,
    }


def test_a_small_gender_group_repeats_the_cards_reason(api, people):
    """The US has four women: no gap, `fewer than 5 women` on each of its rows (D62)."""
    rows = csv_rows(api.get(INSIGHT.format("gender-gap"), {"file": "csv", "country": "US"}))

    assert [row[2:7] for row in rows[1:5]] == [
        ["female", "4", "", "", "fewer than 5 women"],
        ["male", "6", "125000.00", "", "fewer than 5 women"],
        ["other", "0", "", "", "fewer than 5 women"],
        ["undisclosed", "0", "", "", "fewer than 5 women"],
    ]


def test_local_currency_names_the_currency_and_no_rate_month(api, people):
    body = json_body(
        api.get(INSIGHT.format("summary"), {"file": "json", "country": "IN", "currency": "local"})
    )

    assert body["filters"] == {"country": "IN", "on": "2026-10-07", "currency": "INR"}
    assert (body["rows"][0]["currency"], body["rows"][0]["rates"]) == (
        "INR",
        "local currency (INR)",
    )


def test_include_exited_reaches_the_insight_export(api, people):
    body = json_body(api.get(INSIGHT.format("summary"), {"file": "json", "include_exited": "true"}))

    assert body["rows"][0]["headcount"] == 29


def test_on_reaches_the_insight_export(api, people):
    """On 1 January 2027 U1 is paid 130,000 and D5 has joined: 28 people (tests/insights)."""
    body = json_body(api.get(INSIGHT.format("summary"), {"file": "json", "on": "2027-01-01"}))

    assert body["filters"] == {"on": "2027-01-01", "currency": "USD"}
    assert (body["rows"][0]["headcount"], body["rows"][0]["total_cost"]) == (28, "2618450.00")


def test_an_insight_file_is_named_after_the_insight(api, people):
    response = api.get(INSIGHT.format("by-level"), {"file": "xlsx"})

    assert (
        response["Content-Disposition"] == 'attachment; filename="acme-by-level-20261007-0900.xlsx"'
    )
    assert load_workbook(io.BytesIO(response.content)).active.title == "by-level"


@pytest.mark.parametrize("file", ["csv", "pdf"])
def test_an_unknown_insight_is_a_404_before_anything_else(api, lookups, file):
    response = api.get(INSIGHT.format("top-earners"), {"file": file})

    assert (response.status_code, response.json()) == (404, {"detail": "Not found."})


def test_a_bad_insight_filter_is_the_insights_400(api, lookups):
    response = api.get(INSIGHT.format("summary"), {"file": "csv", "country": "XX"})

    assert (response.status_code, response.json()) == (
        400,
        {"country": ["Unknown country_code: XX"]},
    )


def test_an_insight_export_reads_what_the_insight_reads(api, people, django_assert_num_queries):
    """The user, the countries, the departments, the rates, one aggregate, one list of values."""
    with django_assert_num_queries(6):
        response = api.get(INSIGHT.format("by-country"), {"file": "csv"})

    assert response.status_code == 200


def test_a_withheld_trend_point_says_why(api, people):
    """The United Kingdom is G1 alone."""
    rows = csv_rows(api.get(INSIGHT.format("trend"), {"file": "csv", "country": "GB"}))

    assert rows[-1] == ["2026-09-30", "1", "", "fewer than 5 people", "USD", RATES]


def test_trend_exports_both_figures_whatever_its_card_draws(api, people):
    """D62: the measure picks the card's line, and the JSON says which one was asked for."""
    response = api.get(INSIGHT.format("trend"), {"file": "json", "measure": "headcount"})
    body = json_body(response)
    header = csv_rows(api.get(INSIGHT.format("trend"), {"file": "csv"}))[0]

    assert header == ["Month-end", "Headcount", "Total cost", "Reason", "Currency", "Rates"]
    assert body["filters"] == {"measure": "headcount", "on": "2026-10-07", "currency": "USD"}
    assert len(body["rows"]) == 24
    assert body["rows"][-1] == {
        "month_end": "2026-09-30",
        "headcount": 27,
        "total_cost": "2435450.00",
        "reason": None,
        "currency": "USD",
        "rates": RATES,
    }


def test_the_bridge_exports_one_row_per_part(api, people):
    """X1 left on 30 June 2026, the one leaver of DESIGN 9's fixture, so the pay changes show
    their people alone, and the joiners, who are nobody, 0.00. Start is today's 27 and X1's
    132,000 (D62)."""
    rows = csv_rows(api.get(INSIGHT.format("cost-bridge"), {"file": "csv"}))

    assert rows == [
        ["Part", "Date", "People", "Amount", "Reason", "Currency", "Rates"],
        ["Start", "2025-10-07", "28", "2567450.00", "", "USD", RATES],
        ["Leavers", "", "1", "", "fewer than 5 people", "USD", RATES],
        ["Joiners", "", "0", "0.00", "", "USD", RATES],
        ["Pay changes", "", "27", "", "withheld with the leavers", "USD", RATES],
        ["End", "2026-10-07", "27", "2435450.00", "", "USD", RATES],
    ]


def test_by_tenure_exports_its_bands_for_the_level_picked(api, people):
    """D62: level 3 of DESIGN 9's fixture, five people joined six years ago."""
    body = json_body(api.get(INSIGHT.format("by-tenure"), {"file": "json", "level": "3"}))

    assert body["filters"] == {"level": "3", "on": "2026-10-07", "currency": "USD"}
    assert body["rows"][3] == {
        "tenure": "5 to 10 years",
        "headcount": 5,
        "median_gross": "87500.00",
        "reason": None,
        "currency": "USD",
        "rates": RATES,
    }
    assert [row["headcount"] for row in body["rows"]] == [0, 0, 0, 5, 0]


def test_a_withheld_tenure_band_says_why(api, people):
    """Level 6 has three people, I6, I12 and U10."""
    rows = csv_rows(api.get(INSIGHT.format("by-tenure"), {"file": "csv", "level": "6"}))

    assert rows[4] == ["5 to 10 years", "3", "", "fewer than 5 people", "USD", RATES]


def test_cost_mix_exports_one_row_per_country_and_segment(api, people):
    """D62: the take-home view of India, then the United States."""
    query = {"file": "csv", "view": "take_home", "country": "IN,US"}
    rows = csv_rows(api.get(INSIGHT.format("cost-mix"), query))

    assert rows == [
        [
            "Country code",
            "Country",
            "Headcount",
            "Segment",
            "Share %",
            "Reason",
            "Currency",
            "Rates",
        ],
        ["IN", "India", "12", "Net pay", "90.0", "", "USD", RATES],
        ["IN", "India", "12", "Income tax", "10.0", "", "USD", RATES],
        ["US", "United States", "10", "Net pay", "80.0", "", "USD", RATES],
        ["US", "United States", "10", "Income tax", "20.0", "", "USD", RATES],
    ]


def test_a_withheld_country_is_one_row_of_its_headcount(api, people):
    rows = csv_rows(api.get(INSIGHT.format("cost-mix"), {"file": "csv", "country": "GB"}))

    assert rows[1] == ["GB", "United Kingdom", "1", "", "", "fewer than 5 people", "USD", RATES]


@pytest.mark.parametrize("view", ["pay_gap", "bonus", "quartiles", "by_level"])
def test_each_gender_view_exports_its_table_ending_with_reason_currency_and_rates(
    api, people, view
):
    """D60 and D62, for the view the card's control picked."""
    response = api.get(INSIGHT.format("gender-gap"), {"file": "json", "view": view})
    (header, *_) = csv_rows(api.get(INSIGHT.format("gender-gap"), {"file": "csv", "view": view}))

    assert header[-3:] == ["Reason", "Currency", "Rates"]
    assert json_body(response)["filters"]["view"] == view


NL_LOCAL = {"country": "NL", "currency": "local"}
LOCAL = ["EUR", "local currency (EUR)"]


def hire_nl(gender, gross=50_000, *, level=1, bonus=0):
    CompensationRecordFactory(
        employee=EmployeeFactory(country=country("NL"), gender=gender, level=level),
        lines={"BASE": Decimal(gross)},
        bonus_annual=Decimal(bonus),
        effective_from=date(2024, 1, 1),
    )


def test_the_bonus_view_is_one_row_per_group_and_gender_without_a_mean_amount(api, lookups):
    """Five of six women receive 1,000 to 5,000, all six men 2,000 to 7,000: medians 3,000 and
    4,500, and means 3,000 and 4,500, so both gaps are 33.3%. The other employee is counted
    (D62)."""
    for bonus in (1000, 2000, 3000, 4000, 5000, 0):
        hire_nl("female", bonus=bonus)
    for bonus in (2000, 3000, 4000, 5000, 6000, 7000):
        hire_nl("male", bonus=bonus)
    hire_nl("other", bonus=1000)

    rows = csv_rows(
        api.get(INSIGHT.format("gender-gap"), {"file": "csv", "view": "bonus", **NL_LOCAL})
    )

    assert rows[:5] == [
        [
            "Country code",
            "Country",
            "Gender",
            "Headcount",
            "Receiving",
            "Receiving %",
            "Median bonus",
            "Mean gap %",
            "Median gap %",
            "Reason",
            "Currency",
            "Rates",
        ],
        ["NL", "Netherlands", "female", "6", "5", "83.3", "3000.00", "33.3", "33.3", "", *LOCAL],
        ["NL", "Netherlands", "male", "6", "6", "100.0", "4500.00", "33.3", "33.3", "", *LOCAL],
        ["NL", "Netherlands", "other", "1", "", "", "", "33.3", "33.3", "", *LOCAL],
        ["NL", "Netherlands", "undisclosed", "0", "", "", "", "33.3", "33.3", "", *LOCAL],
    ]


QUARTILE_HEADER = [
    "Quartile",
    "Headcount",
    "Women",
    "Men",
    "Women %",
    "Men %",
    "Other not ranked",
    "Undisclosed not ranked",
    "Reason",
    "Currency",
    "Rates",
]


def test_the_quartiles_are_one_row_each_with_the_unranked_beside(api, lookups):
    """Ten women on 50,000 and ten men on 60,000: the women fill the two lower quartiles. The
    other employee is counted on every row and never ranked (D62)."""
    for _ in range(10):
        hire_nl("female", 50_000)
        hire_nl("male", 60_000)
    hire_nl("other", 10_000)

    query = {"file": "csv", "view": "quartiles", **NL_LOCAL}
    rows = csv_rows(api.get(INSIGHT.format("gender-gap"), query))

    assert rows == [
        QUARTILE_HEADER,
        ["1", "5", "5", "0", "100.0", "0.0", "1", "0", "", *LOCAL],
        ["2", "5", "5", "0", "100.0", "0.0", "1", "0", "", *LOCAL],
        ["3", "5", "0", "5", "0.0", "100.0", "1", "0", "", *LOCAL],
        ["4", "5", "0", "5", "0.0", "100.0", "1", "0", "", *LOCAL],
    ]


def test_withheld_quartiles_are_one_row_of_the_headcount(api, people):
    """India's twelve are under twenty."""
    query = {"file": "csv", "view": "quartiles", "country": "IN"}
    rows = csv_rows(api.get(INSIGHT.format("gender-gap"), query))

    assert rows == [
        QUARTILE_HEADER,
        ["", "12", "", "", "", "", "0", "0", "fewer than 20 people", "USD", RATES],
    ]


def test_the_level_gaps_flag_in_the_cards_words_and_count_the_others(api, lookups):
    """Level 1: five women on 40,000 and five men on 50,000, a gap of 20.0% of men's pay,
    flagged in the card's neutral words. Level 2: 50,000 against 51,000, 2.0%, not flagged.
    The other employee is counted on each of the country's rows (D62)."""
    for women, men, level in ((40_000, 50_000, 1), (50_000, 51_000, 2)):
        for _ in range(5):
            hire_nl("female", women, level=level)
            hire_nl("male", men, level=level)
    hire_nl("other", 45_000, level=1)
    query = {"file": "csv", "view": "by_level", **NL_LOCAL}

    rows = csv_rows(api.get(INSIGHT.format("gender-gap"), query))

    assert rows[:4] == [
        [
            "Country code",
            "Country",
            "Level",
            "Women",
            "Women median",
            "Men",
            "Men median",
            "Mean gap %",
            "Median gap %",
            "Flag",
            "Other in country",
            "Undisclosed in country",
            "Reason",
            "Currency",
            "Rates",
        ],
        [
            "NL",
            "Netherlands",
            "1",
            "5",
            "40000.00",
            "5",
            "50000.00",
            "20.0",
            "20.0",
            "5% or more, the threshold in EU Directive 2023/970 Article 10",
            "1",
            "0",
            "",
            *LOCAL,
        ],
        [
            "NL",
            "Netherlands",
            "2",
            "5",
            "50000.00",
            "5",
            "51000.00",
            "2.0",
            "2.0",
            "",
            "1",
            "0",
            "",
            *LOCAL,
        ],
        ["NL", "Netherlands", "3", "0", "", "0", "", "", "", "", "1", "0", "", *LOCAL],
    ]


def json_filters(api, name: str, **query) -> dict:
    return json_body(api.get(INSIGHT.format(name), {"file": "json", **query}))["filters"]


def test_a_file_names_the_control_as_the_insight_applied_it(api, people):
    """D62: the default when none was given, in the list's own spelling when one was."""
    assert json_filters(api, "by-tenure") == {"level": "all", "on": "2026-10-07", "currency": "USD"}
    assert json_filters(api, "gender-gap", view=" BONUS ")["view"] == "bonus"
    assert json_filters(api, "summary", include_exited="true") == {
        "include_exited": "true",
        "on": "2026-10-07",
        "currency": "USD",
    }


@pytest.mark.parametrize("name", ["trend", "cost-bridge"])
def test_a_dated_insights_file_does_not_say_include_exited_applied(api, people, name):
    """D62: trend and cost-bridge count the people employed on each date."""
    filters = json_filters(api, name, include_exited="true", country="US")

    assert "include_exited" not in filters
    assert filters["country"] == "US"
