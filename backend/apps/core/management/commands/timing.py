"""`timing`: measure the directory, insights, import and seed on the seed of 10,000 employees on
staff and 5,000 who left (D64).

Driven by the /perf-timing skill. Step 2 ships the seed case, step 3 the directory cases,
step 5 the import case, step 6 the insight cases and step 7 the export case. Prints a Markdown
row for docs/PERFORMANCE.md plus the per-case breakdown; the export case prints its lines only.
"""

from __future__ import annotations

import io
import os
import platform
import resource
import subprocess
import sys
import tempfile
import time
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.test import Client

from apps.core.dates import REPORTING_START, period_start, today_utc
from apps.insights.queries import CONTROLS
from seed import profiles

CASES = ("directory", "insights", "import", "export", "seed")

# Each directory case is the query string of one request (the /perf-timing skill names them).
DIRECTORY_CASES = {
    "unfiltered": "",
    "country filter": "country=IN",
    "department and gross range": "department=engineering&gross_min=50000&gross_max=150000",
    "sort by reporting gross": "ordering=-gross",
    # Every term is matched against the name, the email, legacy_code and ACME-<id> (D48).
    "search by name": "search=rao",
    "search by code": "search=ACME-17",
    # The directory opens on Active (D64); everyone is one choice away.
    "everyone": "status=any",
}

# Each insight is timed unfiltered and with one country (the /perf-timing skill names them).
INSIGHT_NAMES = (
    "summary",
    "by-country",
    "by-department",
    "distribution",
    "by-level",
    "gender-gap",
    "trend",
    "cost-bridge",
    "by-tenure",
    "cost-mix",
)
INSIGHT_FILTERS = {"unfiltered": "", "country filter": "country=IN"}
# Every other value of a card control is timed as well, since each reads its own figures (D62).
INSIGHT_CONTROLS = tuple(
    (name, f"{key}={value}") for name, (key, values) in CONTROLS.items() for value in values[1:]
)
# The insights whose work grows with the period are timed on each preset and the longest range
# as well (D64); the rest read the state on its last day.
PERIOD_INSIGHTS = ("summary", "trend", "cost-bridge")


def period_cases(today: date) -> dict[str, str]:
    """The dashboard's presets besides the default, and the reporting cut-off to today, as `from`
    and `to` (D64, D70)."""
    fy = date(today.year if (today.month, today.day) >= (4, 1) else today.year - 1, 4, 1)
    return {
        "last 3 months": f"from={period_start(today, 3)}",
        "last 6 months": f"from={period_start(today, 6)}",
        "FY to date": f"from={fy}",
        "last FY": f"from={date(fy.year - 1, 4, 1)}&to={fy - timedelta(days=1)}",
        "since the cut-off": f"from={REPORTING_START}",
    }


# The unfiltered directory is exported once in each format (SPEC F6).
EXPORT_FILES = ("csv", "xlsx", "json")


# ru_maxrss is bytes on macOS and kilobytes on Linux.
RSS_SCALE = 1 if sys.platform == "darwin" else 1024


def git_short_sha() -> str:
    try:
        return subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True, check=True
        ).stdout.strip()
    except OSError, subprocess.CalledProcessError:
        return "unknown"


def machine_label() -> str:
    return f"{platform.system()} {platform.machine()}, {os.cpu_count()} cores"


def percentile(samples: list[float], fraction: float) -> float:
    ordered = sorted(samples)
    return ordered[round(fraction * (len(ordered) - 1))]


def api_client() -> Client:
    """A test client logged in as the seeded HR Manager, so the timings include auth."""
    host = next((h for h in settings.ALLOWED_HOSTS if not h.startswith((".", "*"))), "localhost")
    client = Client(headers={"host": host})
    response = client.post(
        "/api/v1/auth/login/",
        {"email": profiles.HR_MANAGER_EMAIL, "password": settings.SEED_ADMIN_PASSWORD},
        content_type="application/json",
    )
    if response.status_code != 200:
        raise CommandError(
            f"login as {profiles.HR_MANAGER_EMAIL} failed with {response.status_code}; "
            "is the database seeded and SEED_ADMIN_PASSWORD the seeded one?"
        )
    return Client(headers={"host": host, "authorization": f"Bearer {response.json()['access']}"})


def time_requests(client: Client, path: str, repeats: int) -> list[float]:
    """Wall time in ms of `repeats` requests after one unmeasured warm-up."""
    samples: list[float] = []
    for index in range(repeats + 1):
        started = time.perf_counter()
        response = client.get(path)
        elapsed = (time.perf_counter() - started) * 1000
        if response.status_code != 200:
            raise CommandError(f"{path} answered {response.status_code}: {response.content[:300]}")
        if index:
            samples.append(elapsed)
    return samples


def exported_rows(file: str, body: bytes) -> int:
    """Data rows in an export, the header left out."""
    import csv
    import json

    from openpyxl import load_workbook

    if file == "csv":
        return sum(1 for _ in csv.reader(io.StringIO(body.decode("utf-8-sig")))) - 1
    if file == "xlsx":
        # A write-only workbook records no dimension, so the rows are counted.
        sheet = load_workbook(io.BytesIO(body), read_only=True).active
        return sum(1 for _ in sheet.iter_rows(values_only=True)) - 1
    return len(json.loads(body)["rows"])


def time_seed(scratch_db: Path, active: int, exited: int, seed: int) -> float:
    """Wall time of `seed_data --reset` in a fresh process against a scratch database."""
    env = {**os.environ, "DATABASE_PATH": str(scratch_db)}
    manage = str(Path(__file__).resolve().parents[4] / "manage.py")
    subprocess.run([sys.executable, manage, "migrate", "--noinput", "-v", "0"], env=env, check=True)
    started = time.perf_counter()
    subprocess.run(
        [
            sys.executable,
            manage,
            "seed_data",
            "--active",
            str(active),
            "--exited",
            str(exited),
            "--seed",
            str(seed),
            "--reset",
        ],
        env=env,
        check=True,
        capture_output=True,
    )
    return time.perf_counter() - started


class Command(BaseCommand):
    help = "Measure the step targets on the seed and print a PERFORMANCE.md row."

    def add_arguments(self, parser):
        parser.add_argument("--case", choices=("all", *CASES), default="all")
        parser.add_argument("--active", type=int, default=profiles.ACTIVE_EMPLOYEES)
        parser.add_argument("--exited", type=int, default=profiles.EXITED_EMPLOYEES)
        parser.add_argument("--seed", type=int, default=42)
        parser.add_argument("--repeats", type=int, default=20)
        # Each import sheet can run alone, so its peak memory is its own (D64).
        parser.add_argument("--sheet", choices=("both", "all", "new"), default="both")

    def require_seed(self, active: int, exited: int) -> None:
        from apps.employees.models import Employee

        count = Employee.objects.count()
        if count < active + exited:
            raise CommandError(
                f"the database has {count} employees; the timings need the seed of {active} "
                f"active and {exited} exited (python manage.py seed_data --active {active} "
                f"--exited {exited} --seed 42)"
            )

    def time_directory(self, repeats: int) -> str:
        """p50 / p95 of the slowest case; every case is printed."""
        client = api_client()
        worst = (0.0, 0.0)
        for name, query in DIRECTORY_CASES.items():
            samples = time_requests(client, f"/api/v1/employees/?{query}", repeats)
            p50, p95 = percentile(samples, 0.5), percentile(samples, 0.95)
            self.stdout.write(
                f"directory, {name}: p50 {p50:.0f} ms, p95 {p95:.0f} ms ({repeats} requests)"
            )
            if p95 > worst[1]:
                worst = (p50, p95)
        return f"{worst[0]:.0f} / {worst[1]:.0f}"

    def time_insights(self, repeats: int) -> str:
        """p50 / p95 of the slowest insight and filter; every one is printed."""
        client = api_client()
        worst = (0.0, 0.0)
        cases = [(name, "") for name in INSIGHT_NAMES] + list(INSIGHT_CONTROLS)
        for name, control in cases:
            for label, query in INSIGHT_FILTERS.items():
                path = f"/api/v1/insights/{name}/?{'&'.join(q for q in (query, control) if q)}"
                samples = time_requests(client, path, repeats)
                p50, p95 = percentile(samples, 0.5), percentile(samples, 0.95)
                self.stdout.write(
                    f"insight {name}{f' {control}' if control else ''}, {label}: p50 {p50:.0f} ms, "
                    f"p95 {p95:.0f} ms ({repeats} requests)"
                )
                if p95 > worst[1]:
                    worst = (p50, p95)
        for label, query in period_cases(today_utc()).items():
            for name in PERIOD_INSIGHTS:
                samples = time_requests(client, f"/api/v1/insights/{name}/?{query}", repeats)
                p50, p95 = percentile(samples, 0.5), percentile(samples, 0.95)
                self.stdout.write(
                    f"insight {name}, {label}: p50 {p50:.0f} ms, p95 {p95:.0f} ms "
                    f"({repeats} requests)"
                )
                if p95 > worst[1]:
                    worst = (p50, p95)
        return f"{worst[0]:.0f} / {worst[1]:.0f}"

    def time_export(self) -> None:
        """The default directory export, on Active, in each format: rows, size, wall time and
        how far it raised the process's peak memory. SPEC F6 asks for every active employee,
        10,000 on the seed, so a short file is an error (D64)."""
        from apps.employees.models import Employee
        from apps.employees.queries import DEFAULT_STATUS, status_expression

        statuses = Employee.objects.annotate(status=status_expression(today_utc()))
        active = statuses.filter(status=DEFAULT_STATUS).count()

        client = api_client()
        for file in EXPORT_FILES:
            peak_before = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * RSS_SCALE
            started = time.perf_counter()
            response = client.get(f"/api/v1/exports/employees/?file={file}")
            body = b"".join(response.streaming_content) if response.streaming else response.content
            elapsed = time.perf_counter() - started
            peak_after = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * RSS_SCALE
            if response.status_code != 200:
                raise CommandError(f"the {file} export answered {response.status_code}")
            rows = exported_rows(file, body)
            if rows != active:
                raise CommandError(
                    f"the {file} export holds {rows:,} rows, not the {active:,} active employees"
                )
            self.stdout.write(
                f"export, employees as {file}: {rows:,} rows, "
                f"{len(body) / 2**20:.1f} MB in {elapsed:.1f} s, "
                f"peak memory up {(peak_after - peak_before) / 2**20:.0f} MB"
            )

    def time_import(self, seed: int, sheet: str) -> str:
        """Stage and approve the two 15,000-row sheets (D64): `make_import_file --all --raises
        1000`, the whole organisation with a thousand raises, whose times are the row's, and
        `--new 15000`, the heaviest approve. Each runs inside a transaction that is rolled back,
        so the database is left as it was."""
        everyone = ""
        if sheet in ("both", "all"):
            everyone = self.time_sheet("--all --raises 1000", {"all": True, "raises": 1000}, seed)
        if sheet in ("both", "new"):
            self.time_sheet("--new 15000", {"new": 15000}, seed)
        return everyone

    def time_sheet(self, name: str, options: dict, seed: int) -> str:
        """Stage one sheet through the API and approve it, rolled back. Prints both times, how
        far staging raised the process's peak memory and the peak itself."""
        from django.core.management import call_command
        from django.db import transaction

        client = api_client()
        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "import.xlsx"
            call_command(
                "make_import_file", out=str(path), seed=seed, stdout=io.StringIO(), **options
            )
            peak_before = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * RSS_SCALE
            with transaction.atomic(), path.open("rb") as handle:
                started = time.perf_counter()
                response = client.post("/api/v1/imports/", {"file": handle})
                staging = time.perf_counter() - started
                peak_after = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * RSS_SCALE
                if response.status_code != 201:
                    raise CommandError(
                        f"staging answered {response.status_code}: {response.content[:300]}"
                    )
                started = time.perf_counter()
                approved = client.post(f"/api/v1/imports/{response.json()['id']}/approve/")
                approving = time.perf_counter() - started
                transaction.set_rollback(True)
        if approved.status_code != 200:
            raise CommandError(f"approve answered {approved.status_code}: {approved.content[:300]}")
        body, applied = response.json(), approved.json()
        peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * RSS_SCALE
        self.stdout.write(
            f"import {name}, staging {body['total_rows']} rows ({body['valid_rows']} valid, "
            f"{body['warning_rows']} warning, {body['error_rows']} error, {body['noop_rows']} "
            f"no-op): {staging:.1f} s, peak memory up {(peak_after - peak_before) / 2**20:.0f} MB"
        )
        self.stdout.write(
            f"import {name}, approve of {applied['applied_rows']} rows ({applied['revised_rows']} "
            f"revised, {applied['created_rows']} created): {approving:.1f} s, rolled back, "
            f"process peak {peak / 2**20:.0f} MB"
        )
        return f"{staging:.1f} / {approving:.1f}"

    def handle(self, *args, **options):
        wanted = CASES if options["case"] == "all" else (options["case"],)
        results: dict[str, str] = {case: "" for case in CASES}

        if "directory" in wanted:
            self.require_seed(options["active"], options["exited"])
            results["directory"] = self.time_directory(options["repeats"])

        if "insights" in wanted:
            self.require_seed(options["active"], options["exited"])
            results["insights"] = self.time_insights(options["repeats"])

        if "import" in wanted:
            self.require_seed(options["active"], options["exited"])
            results["import"] = self.time_import(options["seed"], options["sheet"])

        if "export" in wanted:
            self.require_seed(options["active"], options["exited"])
            self.time_export()

        if "seed" in wanted:
            with tempfile.TemporaryDirectory() as scratch:
                try:
                    seconds = time_seed(
                        Path(scratch) / "timing.sqlite3",
                        options["active"],
                        options["exited"],
                        options["seed"],
                    )
                except subprocess.CalledProcessError as exc:
                    raise CommandError(f"seed failed: {exc.stderr[-2000:]}") from exc
            results["seed"] = f"{seconds:.1f}"
            self.stdout.write(
                f"seed: {options['active']} active and {options['exited']} exited employees "
                f"with seed {options['seed']} in "
                f"{seconds:.1f} s (fresh process, migrate excluded, --reset included)"
            )

        today = datetime.now(UTC).date().isoformat()
        row = (
            f"| {today} | {git_short_sha()} | {machine_label()} | {results['directory']} | "
            f"{results['insights']} | {results['import']} | {results['seed']} |"
        )
        self.stdout.write(row)
