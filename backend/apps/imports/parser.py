"""Reading an uploaded sheet (SPEC F1, DESIGN 1.5): the file reader and the cell parsers.

The template's shape and the header rules live with every other rule in compensation/validators.py
(SPEC section 7, CLAUDE.md rule 4). Every message text comes from compensation/messages.py, so a
sheet reads the same words as the API.
"""

from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
import zipfile
from collections.abc import Iterator, Mapping, Sequence
from contextlib import closing
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePath
from typing import Any
from xml.parsers import expat

from openpyxl.reader.excel import ExcelReader
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from apps.compensation.messages import msg
from apps.compensation.validators import IDENTITY_COLUMNS, Message, NormalizedRow, error
from apps.core.exceptions import Invalid

DATE_COLUMNS = ("joining_date", "last_date", "effective_from")
# The whole seeded organisation, 10,000 on staff and 5,000 who left, in one file (D64).
ROW_LIMIT = 15_000
# What a sheet may cost before openpyxl reads it (D56). A full 15,000-row template unpacks to
# about 10 MB and is 22 columns wide.
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_ROW_CELLS = 200
EXCEL_MAX_ROW = 1_048_576
# openpyxl builds some parts as a whole tree, at about 100 bytes an element; a full 15,000-row
# template has about 560,000 of them.
MAX_ELEMENTS = 2_000_000
# Cell contents and shared strings openpyxl streams. Anything else, style formats, merged ranges
# and defined names among them, it builds into objects at 560 to 680 bytes each; a full template
# has 339 of those (D56).
CELL_TAGS = frozenset({"row", "c", "v", "f", "is", "t", "r", "si"})
MAX_OBJECT_ELEMENTS = 300_000
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# A browser sends what the operating system calls the file. Windows with Excel installed calls a
# .csv application/vnd.ms-excel, so that type is accepted for .csv too.
CONTENT_TYPES = {
    ".xlsx": {XLSX_TYPE, "application/octet-stream"},
    ".csv": {
        "text/csv",
        "application/csv",
        "text/plain",
        "application/vnd.ms-excel",
        "application/octet-stream",
    },
}

THOUSANDS = re.compile(r"\d{1,3}(?:,\d{3})+")  # 1,200 and 1,200,000
LAKHS = re.compile(r"\d{1,2}(?:,\d{2})*,\d{3}")  # 12,00,000
PLAIN_NUMBER = re.compile(r"\d+(?:\.\d*)?|\.\d+")
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
WHOLE = re.compile(r"\d+(?:\.0+)?")
LEVEL_KEPT_MAX = 999_999  # past this a cell is no level, and parse_level keeps no int of it


# File


@dataclass
class ParsedFile:
    """The header row and the data rows, by their row number in the sheet."""

    header: list[Any]
    rows: list[tuple[int, list[Any]]]
    row_count: int
    # Letters of columns with no header that still hold data: ignored, and said so.
    unnamed: list[str] = field(default_factory=list)


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def check_file_type(file_name: str, content_type: str | None) -> str:
    """The file's extension, once it and the content type agree (SPEC section 8)."""
    extension = PurePath(file_name or "").suffix.lower()
    if extension not in CONTENT_TYPES:
        raise Invalid({"file": [msg("file_extension", name=file_name or "a file with no name")]})
    sent = (content_type or "application/octet-stream").split(";")[0].strip().lower()
    if sent not in CONTENT_TYPES[extension]:
        raise Invalid({"file": [msg("file_content_type", content_type=sent, extension=extension)]})
    return extension


def _binary(upload):
    """The file under a Django upload, or the stream itself in tests."""
    stream = getattr(upload, "file", upload)
    stream.seek(0)
    return stream


def _refused(key: str, number: int, limit: int) -> Invalid:
    return Invalid({"file": [msg(key, row=f"{number:,}", limit=f"{limit:,}")]})


def _too_wide(number: int) -> Invalid:
    return _refused("row_too_wide", number, MAX_ROW_CELLS)


def _forbidden(*_args) -> None:
    # The refusals defusedxml makes: no entity declarations and no external entities, so no
    # entity expansion and no fetching.
    raise Invalid({"file": [msg("file_unreadable")]})


def _row_number(given: str) -> int:
    """A row's `r` as openpyxl reads it: whole, or a float with nothing after the point."""
    try:
        return int(given)
    except ValueError:
        number = float(given)
        if not number.is_integer():
            raise
        return int(number)


class _Scan:
    """Counts the elements of the whole file, and those openpyxl would build into objects. In
    each part it follows the row number and the column of the last cell as openpyxl counts
    them, keeping nothing: a row's `r` or one past the last row, a cell's `r` or one past the
    last cell."""

    def __init__(self) -> None:
        self.elements = 0
        self.objects = 0
        self.row = 0
        self.column = 0

    def start(self, name: str, attributes: dict[str, str]) -> None:
        tag = name.rpartition("}")[2]
        self.elements += 1
        if tag not in CELL_TAGS:
            self.objects += 1
        if self.elements > MAX_ELEMENTS or self.objects > MAX_OBJECT_ELEMENTS:
            raise Invalid({"file": [msg("file_too_big")]})
        if tag == "row":
            given = attributes.get("r")
            self.row, self.column = _row_number(given) if given else self.row + 1, 0
            if self.row > EXCEL_MAX_ROW:
                raise _refused("row_past_limit", self.row, EXCEL_MAX_ROW)
        elif tag == "c":
            reference = attributes.get("r")
            self.column = coordinate_to_tuple(reference)[1] if reference else self.column + 1
            if self.column > MAX_ROW_CELLS:
                raise _too_wide(self.row)


def _scan_part(part, scan: _Scan) -> None:
    scan.row = scan.column = 0
    parser = expat.ParserCreate(namespace_separator="}")
    parser.StartElementHandler = scan.start
    parser.EntityDeclHandler = _forbidden
    parser.UnparsedEntityDeclHandler = _forbidden
    parser.ExternalEntityRefHandler = _forbidden
    try:
        parser.ParseFile(part)
    except expat.ExpatError:
        pass  # not XML, or broken: openpyxl stops at the same byte, after the same rows


def preflight_xlsx(stream) -> None:
    """Refuse a workbook that would cost more to open than the machine has (D56), before
    openpyxl, which reads some parts whole and builds every cell of a row out to the last one:

    - the parts the zip declares, together, over 50 MB unpacked. zipfile stops a part at its
      declared size and fails its checksum, so the declaration holds;
    - a row numbered past Excel's last row, or with a cell past column 200, where openpyxl would
      build every filler row or cell in between;
    - more than 2,000,000 elements in all, which openpyxl would build at 100 bytes each, or
      more than 300,000 that it would build into objects at up to 680 bytes each.

    Every part is scanned, by name or not: openpyxl opens each worksheet at load, and a sheet
    can live under any name. The scan is expat with no tree, so it keeps nothing.
    """
    with zipfile.ZipFile(stream) as archive:
        parts = archive.infolist()
        if sum(part.file_size for part in parts) > MAX_UNPACKED_BYTES:
            raise Invalid({"file": [msg("file_unpacks_large", limit="50 MB")]})
        scan = _Scan()
        for part in parts:
            with archive.open(part) as source:
                _scan_part(source, scan)


class _FirstSheetReader(ExcelReader):
    """openpyxl's own reader, opening the first worksheet only. Its `read_worksheets` sizes every
    sheet the workbook lists and reads every chart sheet's charts, once per entry, so a workbook
    that lists one large part many times would have it read as many times (D57)."""

    def read_worksheets(self) -> None:
        parts = set(self.valid_files)  # a list in openpyxl, searched once per sheet entry
        for sheet, rel in self.parser.find_sheets():
            if rel.target in parts and "chartsheet" not in rel.Type:
                self.wb._sheets.append(
                    ReadOnlyWorksheet(self.wb, sheet.name, rel.target, self.shared_strings)
                )
                return


def open_workbook(stream):
    """The workbook as openpyxl reads it, read-only and values only, with its first worksheet
    alone. External links are never followed: openpyxl reads a linked part once per reference
    to it (D57)."""
    reader = _FirstSheetReader(stream, read_only=True, data_only=True, keep_links=False)
    reader.read()
    return reader.wb


def _xlsx_rows(upload) -> Iterator[list[Any]]:
    stream = _binary(upload)
    # A crafted file makes openpyxl raise a wide range of errors (IndexError on a shared
    # string, ValueError on a cell, ParseError on the XML, among others); each means unreadable.
    try:
        preflight_xlsx(stream)
        stream.seek(0)
        workbook = open_workbook(stream)
    except Invalid:
        raise
    except Exception as exc:
        raise Invalid({"file": [msg("file_unreadable")]}) from exc
    try:
        sheet = workbook.worksheets[0]  # a workbook of chart sheets alone has none
        # The dimension the file declares cannot be trusted, so every row is streamed, up to
        # Excel's own last row.
        sheet.reset_dimensions()
        for values in sheet.iter_rows(max_row=EXCEL_MAX_ROW, values_only=True):
            yield list(values)
    except Exception as exc:
        raise Invalid({"file": [msg("file_unreadable")]}) from exc
    finally:
        workbook.close()


def _csv_rows(upload) -> Iterator[list[Any]]:
    text = io.TextIOWrapper(_binary(upload), encoding="utf-8-sig", newline="")
    try:
        yield from csv.reader(text)
    except UnicodeDecodeError as exc:
        raise Invalid({"file": [msg("not_utf8")]}) from exc
    except csv.Error as exc:
        # A cell longer than csv.field_size_limit(), 131,072 characters, among others.
        raise Invalid({"file": [msg("csv_unreadable")]}) from exc
    finally:
        # Hand the stream back open: Django closes and removes the upload after the request.
        text.detach()


def open_rows(upload, file_name: str, content_type: str | None) -> Iterator[list[Any]]:
    """Every row of the first worksheet, or of the CSV, as a list of cell values."""
    extension = check_file_type(file_name, content_type)
    return _xlsx_rows(upload) if extension == ".xlsx" else _csv_rows(upload)


def read(upload, file_name: str, content_type: str | None) -> ParsedFile:
    """The header and every data row, all-blank rows skipped. Rows are counted while streaming,
    past the limit too, because the sheet's own row count cannot be trusted (SPEC section 8)."""
    # Closed on every way out, so the reader releases the file while it is still open.
    with closing(open_rows(upload, file_name, content_type)) as rows:
        header = next(rows, None)
        if header is None or all(is_blank(cell) for cell in header):
            raise Invalid({"file": [msg("file_no_header")]})
        width = len(header)
        if width > MAX_ROW_CELLS:
            raise _too_wide(1)
        headless = {index for index, cell in enumerate(header) if is_blank(cell)}
        kept: list[tuple[int, list[Any]]] = []
        unnamed: set[int] = set()
        count = 0
        for number, cells in enumerate(rows, start=2):
            # The .xlsx pre-scan has refused wide rows already; a .csv row is checked here.
            if len(cells) > MAX_ROW_CELLS:
                raise _too_wide(number)
            if all(is_blank(cell) for cell in cells):
                continue
            count += 1
            if count <= ROW_LIMIT:
                kept.append((number, list(cells)))
            for index, cell in enumerate(cells):
                if (index in headless or index >= width) and not is_blank(cell):
                    unnamed.add(index)
        if count > ROW_LIMIT:
            raise Invalid({"rows": [msg("row_limit", count=f"{count:,}", limit=f"{ROW_LIMIT:,}")]})
        letters = [get_column_letter(index + 1) for index in sorted(unnamed)]
        return ParsedFile(header=header, rows=kept, row_count=count, unnamed=letters)


# Cells


def parse_number(value: Any, field_name: str) -> Decimal | Message | None:
    """A number cell as Decimal, None when blank, or the message for a value that is not one.

    Currency symbols and spaces are stripped. A comma is a thousands separator only with no dot
    anywhere and in groups that end in three digits (1,200 and 12,00,000); any other comma is the
    decimal-separator error. Decimals are kept as given, so the validators can count them.
    """
    if is_blank(value):
        return None
    invalid = "bonus_invalid" if field_name == "bonus_annual" else "amount_invalid"
    if isinstance(value, bool):
        return error(field_name, invalid, field=field_name, value=str(value).upper())
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return error(field_name, invalid, field=field_name, value=str(value))
        return Decimal(str(value))
    if isinstance(value, Decimal):
        return value
    given = str(value).strip()
    cleaned = "".join(ch for ch in given if not ch.isspace() and unicodedata.category(ch) != "Sc")
    sign = ""
    if cleaned[:1] in ("-", "+"):
        sign, cleaned = cleaned[0], cleaned[1:]
    if "," in cleaned:
        if "." in cleaned or not (THOUSANDS.fullmatch(cleaned) or LAKHS.fullmatch(cleaned)):
            return error(field_name, "comma_decimal", value=given)
        cleaned = cleaned.replace(",", "")
    if not PLAIN_NUMBER.fullmatch(cleaned):
        return error(field_name, invalid, field=field_name, value=given)
    return Decimal(f"{sign}{cleaned}")


def parse_date(value: Any, field_name: str) -> date | Message | None:
    """A date cell, or ISO `YYYY-MM-DD` text; openpyxl gives date cells as datetime."""
    if is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    given = str(value).strip()
    if ISO_DATE.fullmatch(given):
        try:
            return date.fromisoformat(given)
        except ValueError:
            pass
    return error(field_name, "date_invalid", field=field_name, value=given)


def parse_code(value: Any) -> str | None:
    """legacy_code: trimmed text, and a numeric cell as its integer text (58213.0 is 58213)."""
    if is_blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_level(value: Any) -> int | None:
    """A whole number, else None, which the level rule reports. So is a number past 999,999:
    it is no level either, and an int of thousands of digits cannot be written back as text."""
    if is_blank(value) or isinstance(value, bool):
        return None
    if isinstance(value, int):
        level = value
    elif isinstance(value, float):
        level = int(value) if value.is_integer() else None
    else:
        given = str(value).strip()
        # Through Decimal: a float of a few hundred digits is inf, and int(inf) raises.
        level = int(Decimal(given)) if WHOLE.fullmatch(given) else None
    return level if level is not None and level <= LEVEL_KEPT_MAX else None


def parse_text(value: Any) -> str | None:
    if is_blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize(row_number: int, cells: Sequence[Any], mapping: Mapping[int, str]) -> NormalizedRow:
    """One data row as the validators read it (DESIGN 1.5). Parse problems ride along as
    messages, which the structure rules report with the rest."""
    values = {
        column: cells[index] if index < len(cells) else None for index, column in mapping.items()
    }
    problems: list[Message] = []

    def kept(parsed):
        if isinstance(parsed, Message):
            problems.append(parsed)
            return None
        return parsed

    dates = {name: kept(parse_date(values.get(name), name)) for name in DATE_COLUMNS}
    amounts = {
        column: kept(parse_number(value, column))
        for column, value in values.items()
        if column not in IDENTITY_COLUMNS
    }
    email = parse_text(values.get("email"))
    country = parse_text(values.get("country_code"))
    gender = parse_text(values.get("gender"))
    return NormalizedRow(
        row_number=row_number,
        legacy_code=parse_code(values.get("legacy_code")),
        full_name=parse_text(values.get("full_name")),
        email=email.lower() if email else None,
        department=parse_text(values.get("department")),
        designation=parse_text(values.get("designation")),
        level=parse_level(values.get("level")),
        country_code=country.upper() if country else None,
        gender=gender.lower() if gender else None,
        joining_date=dates["joining_date"],
        last_date=dates["last_date"],
        effective_from=dates["effective_from"],
        bonus_annual=kept(parse_number(values.get("bonus_annual"), "bonus_annual")),
        amounts=amounts,
        columns=frozenset(values),
        present=frozenset(column for column, value in values.items() if not is_blank(value)),
        parse_messages=problems,
    )
